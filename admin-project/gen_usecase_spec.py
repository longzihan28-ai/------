from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# 设置默认字体
style = doc.styles['Normal']
font = style.font
font.name = '宋体'
font.size = Pt(11)

# 标题
title = doc.add_heading('表1-1  购买商品用例文档', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 创建表格 (12行2列)
table = doc.add_table(rows=12, cols=2)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# 设置列宽
for row in table.rows:
    row.cells[0].width = Cm(4)
    row.cells[1].width = Cm(12)

# 表格内容
rows_data = [
    ("用例名称", "购买商品"),
    ("用例编号", "UC001"),
    ("参与者", "顾客"),
    ("用例描述", "该用例描述一个系统顾客购买商品的行为"),
    ("触发器", "当顾客选择购买商品时，用例触发"),
    ("前置条件", "顾客需要完成登录操作"),
    ("后置条件", "如果顾客账户余额充足，则生成订单并扣除相应金额；否则提示余额不足，购买失败"),
    ("基本事件流",
     "1. 顾客选择要购买的商品并提交购买请求\n"
     "2. 系统验证顾客身份及购买请求的合法性后做出响应\n"
     "3. 系统验证顾客账户余额是否充足\n"
     "4. 系统统计商品总价格\n"
     "5. 系统从顾客账户余额中扣除相应金额\n"
     "6. 系统生成并保存订单信息\n"
     "7. 系统生成购买确认页面并发送给顾客"),
    ("扩展事件流",
     "A-3 如果顾客账户余额不足，系统提示余额不足，取消购买并终止用例\n"
     "A-4 如果顾客拥有优惠券并选择使用，系统验证优惠券有效性后抵扣部分金额，重新计算支付金额"),
    ("结论", "当顾客收到系统发送的购买确认页面或余额不足等异常信息时，用例结束"),
    ("数据需求",
     "D-1 订单信息包括订单号、顾客账号、商品名称、商品数量、商品单价、支付总金额、优惠券抵扣金额"),
    ("业务规则",
     "B-1 只有当顾客账户余额大于等于商品总价格时才能完成购买\n"
     "B-2 优惠券抵扣金额不得超过商品总价格"),
]

for i, (label, content) in enumerate(rows_data):
    cell_label = table.cell(i, 0)
    cell_content = table.cell(i, 1)

    # 设置标签列
    cell_label.text = label
    for p in cell_label.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.name = '宋体'

    # 设置内容列
    cell_content.text = content
    for p in cell_content.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.size = Pt(11)
            run.font.name = '宋体'

# 保存到桌面
output_path = r"c:\Users\Lenovo\Desktop\购买商品用例文档.docx"
doc.save(output_path)
print(f"已生成: {output_path}")
